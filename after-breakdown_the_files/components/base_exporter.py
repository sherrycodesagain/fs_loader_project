import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from config import COMPANY_NAME, DENOMINATION, CURRENCY
from cache.memory_store import memory_store

DERIVED_LABELS = [
    "Sales",
    "Cost of sales", 
    "Operating expenses",
    "Other income/expenses",
    "Income taxes"
]

class BaseExporter:
    def __init__(self, enriched_df: pd.DataFrame, value_date: pd.Timestamp,
                 fiscal_year_end: pd.Timestamp, output_path: str = "output/income_statement.xlsx"):
        self.df = enriched_df.copy()
        self.value_date = value_date
        self.fiscal_year_end = fiscal_year_end
        self.output_path = output_path
        self.wb = Workbook()
        self.row_tracker = {} 

        # Default styles
        self.arial_font = Font(name="Arial", size=10)
        self.bold_font = Font(name="Arial", bold=True, size=10)
        self.italic_font = Font(name="Arial", italic=True, size=10)
        self.level4_font = Font(name="Arial", size=8.5)
        self.header_fill = PatternFill("solid", fgColor="DDEBF7")
        self.total_fill = PatternFill("solid", fgColor="F2F2F2")
        self.thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        self.currency_format = '#,##0_);(#,##0)'
        self.percent_format = '0%;(0%)'

        # To be built inside methods
        self.annual_pivot = None
        self.quarterly_pivot = None
        self.monthly_pivot = None

    def _get_col_idx(self, i: int, mode: str = "annual") -> int:
        return 2 + i * 2 if mode == "annual" else 2 + i