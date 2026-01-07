from openpyxl.utils import get_column_letter

class StatementWriter:
    def __init__(self, base_exporter, row_writer, derived_calc, formatter, delta_calc):
        self.base = base_exporter
        self.row_writer = row_writer
        self.derived_calc = derived_calc
        self.formatter = formatter
        self.delta_calc = delta_calc

    def write_generic_income_statement(self, sheet_title: str, pivot_df, header_title: str, mode: str = "annual"):
        ws = self.base.wb.create_sheet(title=sheet_title)
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.sheet_properties.outlinePr.applyStyles = True
        ws.sheet_properties.outlinePr.showOutlineSymbols = True

        print(f"[INFO] Writing {sheet_title} sheet...")

        # Temporarily swap pivot
        original_pivot = self.base.annual_pivot
        self.base.annual_pivot = pivot_df

        fiscal_columns = list(pivot_df.columns)
        start_row = 7
        is_df = self.base.df[self.base.df["type"] == "IS"]
        top_levels = is_df["top_level_grouping"].dropna().unique()
        current_row = start_row
        row_tracker = {}

        if mode == "annual":
            self.formatter.preallocate_percentage_columns(ws, fiscal_columns)

        for top in top_levels:
            current_row = self.row_writer.write_top_level_grouping(
                ws, is_df, top, fiscal_columns, current_row, row_tracker, mode
            )
            top_key = top.strip().lower()
            current_row += 1

            if top_key == "cost of sales":
                current_row = self.derived_calc.insert_derived_row(
                    ws, "Gross Margin", fiscal_columns,
                    formula=lambda c: f"={get_column_letter(c)}{row_tracker['sales']}-{get_column_letter(c)}{row_tracker['cost of sales']}",
                    base_row=current_row,
                    row_tracker=row_tracker,
                    mode=mode,
                    sales_row=row_tracker.get("sales")
                )
                current_row += 1

            if top_key == "operating expenses":
                gross_row = row_tracker.get("gross margin")
                op_total_row = row_tracker.get("operating expenses")
                current_row = self.derived_calc.insert_derived_row(
                    ws, "Operating Margin", fiscal_columns,
                    formula=lambda c: f"={get_column_letter(c)}{gross_row}-{get_column_letter(c)}{op_total_row}",
                    base_row=current_row,
                    row_tracker=row_tracker,
                    mode=mode,
                    sales_row=row_tracker.get("sales")
                )
                current_row += 1

            if top_key == "other income/expenses":
                op_margin_row = row_tracker.get("operating margin")
                other_total_row = row_tracker.get("other income/expenses")
                current_row = self.derived_calc.insert_derived_row(
                    ws, "Net Income Before Taxes", fiscal_columns,
                    formula=lambda c: f"={get_column_letter(c)}{op_margin_row}-{get_column_letter(c)}{other_total_row}",
                    base_row=current_row,
                    row_tracker=row_tracker,
                    mode=mode,
                    sales_row=row_tracker.get("sales")
                )
                current_row += 1

            if top_key == "income taxes":
                before_tax_row = row_tracker.get("net income before taxes")
                tax_row = row_tracker.get("income taxes")
                current_row = self.derived_calc.insert_derived_row(
                    ws, "Net Income", fiscal_columns,
                    formula=lambda c: f"={get_column_letter(c)}{before_tax_row}-{get_column_letter(c)}{tax_row}",
                    base_row=current_row,
                    row_tracker=row_tracker,
                    mode=mode,
                    sales_row=row_tracker.get("sales")
                )
                current_row += 1

                # Add EBITDA Subcomponents and Total
                current_row, _ = self.derived_calc.write_ebitda_subcomponents(ws, row_tracker, current_row, len(fiscal_columns), mode)
                net_income_row = row_tracker.get("net income")
                ebitda_adj_row = current_row - 1
                current_row = self.derived_calc.insert_derived_row(
                    ws, "EBITDA", fiscal_columns,
                    formula=lambda c: f"={get_column_letter(c)}{net_income_row}+{get_column_letter(c)}{ebitda_adj_row}",
                    base_row=current_row,
                    row_tracker=row_tracker,
                    mode=mode,
                    sales_row=row_tracker.get("sales")
                )
                current_row += 1

        if mode == "annual":
            self.formatter.fill_percentage_formulas(ws, fiscal_columns, start_row)
        if mode in ["quarterly", "monthly"]:
            self.formatter.apply_column_widths(ws, mode, len(fiscal_columns))

        self.formatter.write_header(ws, fiscal_columns, mode)

        if mode == "annual":
            self.delta_calc.insert_delta_columns(ws, fiscal_columns, start_row)

        self.formatter.bold_top_level_rows(ws, start_row)

        for row in range(ws.min_row, ws.max_row + 1):
            if ws.row_dimensions[row].outlineLevel > 0:
                ws.row_dimensions[row].hidden = True
                ws.row_dimensions[row].collapsed = True

        # Restore original pivot
        self.base.annual_pivot = original_pivot