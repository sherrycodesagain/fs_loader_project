# from openpyxl.utils import get_column_letter

# class StatementWriter:
#     def __init__(self, base_exporter, row_writer, derived_calc, formatter, delta_calc):
#         self.base = base_exporter
#         self.row_writer = row_writer
#         self.derived_calc = derived_calc
#         self.formatter = formatter
#         self.delta_calc = delta_calc

#     def write_generic_income_statement(self, sheet_title: str, pivot_df, header_title: str, mode: str = "annual"):
#         ws = self.base.wb.create_sheet(title=sheet_title)
#         ws.sheet_properties.outlinePr.summaryBelow = False
#         ws.sheet_properties.outlinePr.applyStyles = True
#         ws.sheet_properties.outlinePr.showOutlineSymbols = True

#         print(f"[INFO] Writing {sheet_title} sheet...")

#         # Temporarily swap pivot
#         original_pivot = self.base.annual_pivot
#         self.base.annual_pivot = pivot_df

#         fiscal_columns = list(pivot_df.columns)
#         start_row = 7
#         is_df = self.base.df[self.base.df["type"] == "IS"]
#         top_levels = is_df["top_level_grouping"].dropna().unique()
#         current_row = start_row
#         row_tracker = {}

#         if mode == "annual":
#             self.formatter.preallocate_percentage_columns(ws, fiscal_columns)

#         for top in top_levels:
#             current_row = self.row_writer.write_top_level_grouping(
#                 ws, is_df, top, fiscal_columns, current_row, row_tracker, mode
#             )
#             top_key = top.strip().lower()
#             current_row += 1

#             if top_key == "cost of sales":
#                 current_row = self.derived_calc.insert_derived_row(
#                     ws, "Gross Margin", fiscal_columns,
#                     formula=lambda c: f"={get_column_letter(c)}{row_tracker['sales']}-{get_column_letter(c)}{row_tracker['cost of sales']}",
#                     base_row=current_row,
#                     row_tracker=row_tracker,
#                     mode=mode,
#                     sales_row=row_tracker.get("sales")
#                 )
#                 current_row += 1

#             if top_key == "operating expenses":
#                 gross_row = row_tracker.get("gross margin")
#                 op_total_row = row_tracker.get("operating expenses")
#                 current_row = self.derived_calc.insert_derived_row(
#                     ws, "Operating Margin", fiscal_columns,
#                     formula=lambda c: f"={get_column_letter(c)}{gross_row}-{get_column_letter(c)}{op_total_row}",
#                     base_row=current_row,
#                     row_tracker=row_tracker,
#                     mode=mode,
#                     sales_row=row_tracker.get("sales")
#                 )
#                 current_row += 1

#             if top_key == "other income/expenses":
#                 op_margin_row = row_tracker.get("operating margin")
#                 other_total_row = row_tracker.get("other income/expenses")
#                 current_row = self.derived_calc.insert_derived_row(
#                     ws, "Net Income Before Taxes", fiscal_columns,
#                     formula=lambda c: f"={get_column_letter(c)}{op_margin_row}-{get_column_letter(c)}{other_total_row}",
#                     base_row=current_row,
#                     row_tracker=row_tracker,
#                     mode=mode,
#                     sales_row=row_tracker.get("sales")
#                 )
#                 current_row += 1

#             if top_key == "income taxes":
#                 before_tax_row = row_tracker.get("net income before taxes")
#                 tax_row = row_tracker.get("income taxes")
#                 current_row = self.derived_calc.insert_derived_row(
#                     ws, "Net Income", fiscal_columns,
#                     formula=lambda c: f"={get_column_letter(c)}{before_tax_row}-{get_column_letter(c)}{tax_row}",
#                     base_row=current_row,
#                     row_tracker=row_tracker,
#                     mode=mode,
#                     sales_row=row_tracker.get("sales")
#                 )
#                 current_row += 1

#                 # Add EBITDA Subcomponents and Total
#                 current_row, _ = self.derived_calc.write_ebitda_subcomponents(ws, row_tracker, current_row, len(fiscal_columns), mode)
#                 net_income_row = row_tracker.get("net income")
#                 ebitda_adj_row = current_row - 1
#                 current_row = self.derived_calc.insert_derived_row(
#                     ws, "EBITDA", fiscal_columns,
#                     formula=lambda c: f"={get_column_letter(c)}{net_income_row}+{get_column_letter(c)}{ebitda_adj_row}",
#                     base_row=current_row,
#                     row_tracker=row_tracker,
#                     mode=mode,
#                     sales_row=row_tracker.get("sales")
#                 )
#                 current_row += 1

#         if mode == "annual":
#             self.formatter.fill_percentage_formulas(ws, fiscal_columns, start_row)
#         if mode in ["quarterly", "monthly"]:
#             self.formatter.apply_column_widths(ws, mode, len(fiscal_columns))

#         self.formatter.write_header(ws, fiscal_columns, mode)

#         if mode == "annual":
#             self.delta_calc.insert_delta_columns(ws, fiscal_columns, start_row)

#         self.formatter.bold_top_level_rows(ws, start_row)

#         for row in range(ws.min_row, ws.max_row + 1):
#             if ws.row_dimensions[row].outlineLevel > 0:
#                 ws.row_dimensions[row].hidden = True
#                 ws.row_dimensions[row].collapsed = True

#         # Restore original pivot
#         self.base.annual_pivot = original_pivot




from contextlib import contextmanager
from openpyxl.utils import get_column_letter


@contextmanager
def temporary_pivot(base, pivot_df):
    """
    Temporarily swap the base.annual_pivot with a new pivot_df.
    Restores the original pivot even if an error occurs.
    """
    original_pivot = base.annual_pivot
    base.annual_pivot = pivot_df
    try:
        yield
    finally:
        base.annual_pivot = original_pivot


class StatementWriter:
    def __init__(self, base_exporter, row_writer, derived_calc, formatter, delta_calc):
        self.base = base_exporter
        self.row_writer = row_writer
        self.derived_calc = derived_calc
        self.formatter = formatter
        self.delta_calc = delta_calc

    def write_generic_income_statement(self, sheet_title: str, pivot_df, header_title: str, mode: str = "annual"):
        ws = self.base.wb.create_sheet(title=sheet_title)
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.sheet_properties.outlinePr.applyStyles = True
        ws.sheet_properties.outlinePr.showOutlineSymbols = True

        print(f"[INFO] Writing {sheet_title} sheet...")

        # Prepare frequently used values
        fiscal_columns = list(pivot_df.columns)
        col_letters = [get_column_letter(i) for i in range(1, len(fiscal_columns) + 1)]
        start_row = 7

        # Filter IS dataframe once
        is_df = self.base.df.loc[self.base.df["type"] == "IS"]
        top_levels = is_df.loc[is_df["top_level_grouping"].notna(), "top_level_grouping"].unique()

        current_row = start_row
        row_tracker = {}

        if mode == "annual":
            self.formatter.preallocate_percentage_columns(ws, fiscal_columns)

        # Use context manager to safely swap pivot
        with temporary_pivot(self.base, pivot_df):
            for top in top_levels:
                current_row = self.row_writer.write_top_level_grouping(
                    ws, is_df, top, fiscal_columns, current_row, row_tracker, mode
                )
                top_key = top.strip().lower()
                current_row += 1

                # Cost of Sales → Gross Margin
                if top_key == "cost of sales":
                    sales_row = row_tracker["sales"]
                    cos_row = row_tracker["cost of sales"]
                    current_row = self.derived_calc.insert_derived_row(
                        ws, "Gross Margin", fiscal_columns,
                        formula=lambda c: f"={col_letters[c-1]}{sales_row}-{col_letters[c-1]}{cos_row}",
                        base_row=current_row, row_tracker=row_tracker, mode=mode, sales_row=sales_row
                    )
                    current_row += 1

                # Operating Expenses → Operating Margin
                if top_key == "operating expenses":
                    gross_row = row_tracker["gross margin"]
                    op_total_row = row_tracker["operating expenses"]
                    sales_row = row_tracker["sales"]
                    current_row = self.derived_calc.insert_derived_row(
                        ws, "Operating Margin", fiscal_columns,
                        formula=lambda c: f"={col_letters[c-1]}{gross_row}-{col_letters[c-1]}{op_total_row}",
                        base_row=current_row, row_tracker=row_tracker, mode=mode, sales_row=sales_row
                    )
                    current_row += 1

                # Other Income/Expenses → Net Income Before Taxes
                if top_key == "other income/expenses":
                    op_margin_row = row_tracker["operating margin"]
                    other_total_row = row_tracker["other income/expenses"]
                    sales_row = row_tracker["sales"]
                    current_row = self.derived_calc.insert_derived_row(
                        ws, "Net Income Before Taxes", fiscal_columns,
                        formula=lambda c: f"={col_letters[c-1]}{op_margin_row}-{col_letters[c-1]}{other_total_row}",
                        base_row=current_row, row_tracker=row_tracker, mode=mode, sales_row=sales_row
                    )
                    current_row += 1

                # Income Taxes → Net Income + EBITDA
                if top_key == "income taxes":
                    before_tax_row = row_tracker["net income before taxes"]
                    tax_row = row_tracker["income taxes"]
                    sales_row = row_tracker["sales"]

                    # Net Income
                    current_row = self.derived_calc.insert_derived_row(
                        ws, "Net Income", fiscal_columns,
                        formula=lambda c: f"={col_letters[c-1]}{before_tax_row}-{col_letters[c-1]}{tax_row}",
                        base_row=current_row, row_tracker=row_tracker, mode=mode, sales_row=sales_row
                    )
                    current_row += 1

                    # EBITDA subcomponents and EBITDA
                    current_row, _ = self.derived_calc.write_ebitda_subcomponents(
                        ws, row_tracker, current_row, len(fiscal_columns), mode
                    )
                    net_income_row = row_tracker["net income"]
                    ebitda_adj_row = current_row - 1
                    current_row = self.derived_calc.insert_derived_row(
                        ws, "EBITDA", fiscal_columns,
                        formula=lambda c: f"={col_letters[c-1]}{net_income_row}+{col_letters[c-1]}{ebitda_adj_row}",
                        base_row=current_row, row_tracker=row_tracker, mode=mode, sales_row=sales_row
                    )
                    current_row += 1

        # Final formatting and calculations
        if mode == "annual":
            self.formatter.fill_percentage_formulas(ws, fiscal_columns, start_row)
        elif mode in ("quarterly", "monthly"):
            self.formatter.apply_column_widths(ws, mode, len(fiscal_columns))

        self.formatter.write_header(ws, fiscal_columns, mode)

        if mode == "annual":
            self.delta_calc.insert_delta_columns(ws, fiscal_columns, start_row)

        self.formatter.bold_top_level_rows(ws, start_row)

        # Hide/collapse rows with outline level > 0
        for _, dim in ws.row_dimensions.items():
            if dim.outlineLevel > 0:
                dim.hidden = True
                dim.collapsed = True
