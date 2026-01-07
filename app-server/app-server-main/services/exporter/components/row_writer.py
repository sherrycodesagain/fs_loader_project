from openpyxl.styles import Alignment
from .base_exporter import DERIVED_LABELS

class RowWriter:
    def __init__(self, base_exporter):
        self.base = base_exporter

    def write_top_level_grouping(self, ws, is_df, top, fiscal_columns, current_row, row_tracker, mode):
        ws.cell(row=current_row, column=1, value=top)
        ws.cell(row=current_row, column=1).font = self.base.arial_font
        ws.cell(row=current_row, column=1).alignment = Alignment(indent=0)
        ws.row_dimensions[current_row].outlineLevel = 0
        ws.row_dimensions[current_row].hidden = False

        for idx, col in enumerate(fiscal_columns):
            keys = [
                (fs, detail, acct)
                for fs in is_df[is_df["top_level_grouping"] == top]["fs_grouping"].dropna().unique()
                for detail in is_df[(is_df["top_level_grouping"] == top) & (is_df["fs_grouping"] == fs)]["detailed_grouping"].dropna().unique()
                for acct in is_df[(is_df["top_level_grouping"] == top) & (is_df["fs_grouping"] == fs) & (is_df["detailed_grouping"] == detail)]["account"].unique()
                if (fs, detail, acct) in self.base.annual_pivot.index
            ]
            val = sum(self.base.annual_pivot.at[k, col] for k in keys)
            col_idx = self.base._get_col_idx(idx, mode)
            ws.cell(row=current_row, column=col_idx, value=val).number_format = self.base.currency_format

        ws.row_dimensions[current_row].outlineLevel = 0
        ws.row_dimensions[current_row].hidden = False

        fs_groupings = sorted(is_df[is_df["top_level_grouping"] == top]["fs_grouping"].dropna().unique())

        if top.strip().lower() in [label.lower() for label in DERIVED_LABELS]:
            row_tracker[top.strip().lower()] = current_row

        current_row += 1

        for fs in fs_groupings:
            current_row = self.write_fs_grouping(
                ws, is_df, top, fs, fiscal_columns, current_row, row_tracker, mode
            )

        return current_row
    
    # def write_fs_grouping(self, ws, is_df, top, fs, fiscal_columns, current_row, row_tracker, mode):
    #     ws.cell(row=current_row, column=1, value=fs)
    #     ws.cell(row=current_row, column=1).font = self.base.arial_font
    #     ws.cell(row=current_row, column=1).alignment = Alignment(indent=1)
    #     ws.row_dimensions[current_row].outlineLevel = 1
    #     ws.row_dimensions[current_row].hidden = False

    #     keys = [
    #         (fs, detail, acct)
    #         for detail in is_df[(is_df["top_level_grouping"] == top) & (is_df["fs_grouping"] == fs)]["detailed_grouping"].dropna().unique()
    #         for acct in is_df[(is_df["top_level_grouping"] == top) & (is_df["fs_grouping"] == fs) & (is_df["detailed_grouping"] == detail)]["account"].unique()
    #         if (fs, detail, acct) in self.base.annual_pivot.index
    #     ]
    #     for idx, col in enumerate(fiscal_columns):
    #         subtotal = sum(self.base.annual_pivot.at[k, col] for k in keys)
    #         col_idx = self.base._get_col_idx(idx, mode)
    #         ws.cell(row=current_row, column=col_idx, value=subtotal).number_format = self.base.currency_format

    #     fs_clean = fs.strip().lower()
    #     if fs_clean in ["amortization", "interest expense"]:
    #         row_tracker[fs_clean] = current_row

    #     current_row += 1

    #     detail_groups = sorted(
    #         is_df[(is_df["top_level_grouping"] == top) & (is_df["fs_grouping"] == fs)]["detailed_grouping"].dropna().unique()
    #     )

    #     for detail in detail_groups:
    #         current_row = self.write_detailed_grouping(
    #             ws, is_df, top, fs, detail, fiscal_columns, current_row, row_tracker, mode
    #         )

    #     return current_row

#improved below
    def write_fs_grouping(self, ws, is_df, top, fs, fiscal_columns, current_row, row_tracker, mode):
        ws.cell(row=current_row, column=1, value=fs)
        ws.cell(row=current_row, column=1).font = self.base.arial_font
        ws.cell(row=current_row, column=1).alignment = Alignment(indent=1)
        ws.row_dimensions[current_row].outlineLevel = 1
        ws.row_dimensions[current_row].hidden = False

        # Pre-filter once
        top_fs_mask = (is_df["top_level_grouping"] == top) & (is_df["fs_grouping"] == fs)
        details = is_df.loc[top_fs_mask, "detailed_grouping"].dropna().unique()

        # Build keys and get values in one go
        keys = [
            (fs, detail, acct)
            for detail in details
            for acct in is_df.loc[top_fs_mask & (is_df["detailed_grouping"] == detail), "account"].unique()
            if (fs, detail, acct) in self.base.annual_pivot.index
        ]

        if keys:
            pivot_slice = self.base.annual_pivot.loc[keys]
            for idx, col in enumerate(fiscal_columns):
                subtotal = pivot_slice[col].sum()
                col_idx = self.base._get_col_idx(idx, mode)
                ws.cell(row=current_row, column=col_idx, value=subtotal).number_format = self.base.currency_format

        fs_clean = fs.strip().lower()
        if fs_clean in ["amortization", "interest expense"]:
            row_tracker[fs_clean] = current_row

        current_row += 1

        # Sort details once
        detail_groups = sorted(details)
        for detail in detail_groups:
            current_row = self.write_detailed_grouping(
                ws, is_df, top, fs, detail, fiscal_columns, current_row, row_tracker, mode
            )

        return current_row



    def write_detailed_grouping(self, ws, is_df, top, fs, detail, fiscal_columns, current_row, row_tracker, mode):
        detail_row = current_row
        ws.cell(row=detail_row, column=1, value=detail)
        ws.cell(row=detail_row, column=1).font = self.base.arial_font
        ws.cell(row=detail_row, column=1).alignment = Alignment(indent=2)
        ws.row_dimensions[detail_row].outlineLevel = 2
        ws.row_dimensions[detail_row].hidden = False

        keys = [
            (fs, detail, acct)
            for acct in is_df[
                (is_df["top_level_grouping"] == top)
                & (is_df["fs_grouping"] == fs)
                & (is_df["detailed_grouping"] == detail)
            ]["account"].unique()
            if (fs, detail, acct) in self.base.annual_pivot.index
        ]

        for idx, col in enumerate(fiscal_columns):
            subtotal = sum(self.base.annual_pivot.at[k, col] for k in keys)
            col_idx = self.base._get_col_idx(idx, mode)
            ws.cell(row=current_row, column=col_idx, value=subtotal).number_format = self.base.currency_format

        current_row += 1

        # Write account-level rows under this detailed group
        current_row = self.write_account_rows(
            ws, is_df, top, fs, detail, fiscal_columns, current_row, row_tracker, mode
        )

        # Group account rows under the detail
        if current_row > detail_row + 1:
            ws.row_dimensions.group(detail_row + 1, current_row - 1, outline_level=3, hidden=True)

        return current_row

    def write_account_rows(self, ws, is_df, top, fs, detail, fiscal_columns, current_row, row_tracker, mode):
        account_rows = is_df[
            (is_df["top_level_grouping"] == top) &
            (is_df["fs_grouping"] == fs) &
            (is_df["detailed_grouping"] == detail)
        ]

        for acct in sorted(account_rows["account"].unique()):
            ws.cell(row=current_row, column=1, value=acct)
            ws.cell(row=current_row, column=1).font = self.base.level4_font
            ws.cell(row=current_row, column=1).alignment = Alignment(indent=3)

            for idx, col in enumerate(fiscal_columns):
                value = self.base.annual_pivot.at[(fs, detail, acct), col] if (fs, detail, acct) in self.base.annual_pivot.index else 0
                col_idx = self.base._get_col_idx(idx, mode)
                ws.cell(row=current_row, column=col_idx, value=value).number_format = self.base.currency_format

            ws.row_dimensions[current_row].outlineLevel = 3
            ws.row_dimensions[current_row].hidden = True
            current_row += 1

        return current_row