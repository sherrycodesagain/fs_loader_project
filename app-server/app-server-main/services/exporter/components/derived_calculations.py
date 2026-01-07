from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment

class DerivedCalculations:
    def __init__(self, base_exporter):
        self.base = base_exporter

    def insert_derived_row(self, ws, label, fiscal_columns, formula, base_row, row_tracker, mode, sales_row=None):
        ws.cell(row=base_row, column=1, value=label)
        ws.cell(row=base_row, column=1).font = self.base.bold_font
        ws.cell(row=base_row, column=1).alignment = Alignment(indent=0)

        for i, _ in enumerate(fiscal_columns):
            col_idx = self.base._get_col_idx(i, mode)
            formula_str = formula(col_idx)
            ws.cell(row=base_row, column=col_idx, value=formula_str)
            ws.cell(row=base_row, column=col_idx).number_format = self.base.currency_format
            ws.cell(row=base_row, column=col_idx).font = self.base.bold_font

        row_tracker[label.strip().lower()] = base_row
        next_row = base_row + 1

        # ⬇️ For non-annual: add inline % row
        if mode != "annual" and sales_row:
            self.insert_percentage_below_row(ws, base_row, label, sales_row, fiscal_columns, mode)
            next_row += 1

        return next_row

    def insert_percentage_below_row(self, ws, label_row, label_text, sales_row, fiscal_columns, mode):
        pct_row = label_row + 1
        ws.insert_rows(pct_row)

        ws.cell(row=pct_row, column=1, value=f"{label_text} (%)")
        ws.cell(row=pct_row, column=1).font = self.base.italic_font
        ws.cell(row=pct_row, column=1).alignment = Alignment(indent=0)

        for i, _ in enumerate(fiscal_columns):
            col_idx = self.base._get_col_idx(i, mode)

            letter = get_column_letter(col_idx)
            formula = f"=IF(${letter}${sales_row}=0, 0, {letter}{label_row}/${letter}${sales_row})"

            cell = ws.cell(row=pct_row, column=col_idx)
            cell.value = formula
            cell.number_format = self.base.percent_format
            cell.font = self.base.arial_font

    # def write_ebitda_subcomponents(self, ws, row_tracker, start_row, num_fiscal, mode):
    #     """
    #     Write EBITDA components (excluding Net Income) as level 2 rows,
    #     and return a dict mapping fiscal column index → row numbers of the subcomponents.
    #     """
    #     subcomponents = ["Amortization", "Interest expense", "Income taxes"]
    #     subcomponent_rows = {}  

    #     for i in range(num_fiscal):
    #         col_idx = self.base._get_col_idx(i, mode)
    #         if col_idx not in subcomponent_rows:
    #             subcomponent_rows[col_idx] = []

    #     for label in subcomponents:
    #         key = label.strip().lower()
    #         source_row = row_tracker.get(key)
    #         if not source_row:
    #             continue

    #         ws.cell(row=start_row, column=1, value=label)
    #         ws.cell(row=start_row, column=1).font = self.base.arial_font
    #         ws.cell(row=start_row, column=1).alignment = Alignment(indent=1)

    #         for i in range(num_fiscal):
    #             col_idx = self.base._get_col_idx(i, mode)
    #             pct_col_idx = col_idx + 1

    #             if col_idx not in subcomponent_rows:
    #                 subcomponent_rows[col_idx] = []
    #             subcomponent_rows[col_idx].append(start_row)

    #             # Reference original value row (source_row)
    #             ws.cell(row=start_row, column=col_idx).value = f"={get_column_letter(col_idx)}{source_row}"
    #             ws.cell(row=start_row, column=col_idx).number_format = self.base.currency_format
    #             ws.cell(row=start_row, column=col_idx).font = self.base.arial_font

    #             # % of Sales
    #             sales_row = row_tracker.get("sales")
    #             if sales_row:
    #                 ws.cell(row=start_row, column=pct_col_idx).value = (
    #                     f"={get_column_letter(col_idx)}{start_row}/${get_column_letter(col_idx)}${sales_row}"
    #                 )
    #                 ws.cell(row=start_row, column=pct_col_idx).number_format = self.base.percent_format
    #                 ws.cell(row=start_row, column=pct_col_idx).font = self.base.arial_font

    #         ws.row_dimensions[start_row].outlineLevel = 2
    #         ws.row_dimensions[start_row].hidden = True
    #         start_row += 1

    #     subtotal_row = start_row
    #     ws.cell(row=subtotal_row, column=1, value="total")
    #     ws.cell(row=subtotal_row, column=1).font = self.base.italic_font
    #     ws.cell(row=subtotal_row, column=1).alignment = Alignment(indent=1)

    #     for i in range(num_fiscal):
    #         col_idx = 2 + i * 2
    #         pct_col_idx = col_idx + 1

    #         row_nums = subcomponent_rows.get(col_idx, [])
    #         if row_nums:
    #             formula = "+".join([f"{get_column_letter(col_idx)}{r}" for r in row_nums])
    #             formula = f"={formula}"
    #         else:
    #             formula = "=0"

    #         ws.cell(row=subtotal_row, column=col_idx).value = formula
    #         ws.cell(row=subtotal_row, column=col_idx).number_format = self.base.currency_format
    #         ws.cell(row=subtotal_row, column=col_idx).font = self.base.bold_font

    #         sales_row = row_tracker.get("sales")
    #         if sales_row:
    #             ws.cell(row=subtotal_row, column=pct_col_idx).value = (
    #                 f"={get_column_letter(col_idx)}{subtotal_row}/${get_column_letter(col_idx)}${sales_row}"
    #             )
    #         else:
    #             ws.cell(row=subtotal_row, column=pct_col_idx).value = 0

    #         ws.cell(row=subtotal_row, column=pct_col_idx).number_format = self.base.percent_format
    #         ws.cell(row=subtotal_row, column=pct_col_idx).font = self.base.bold_font

    #     ws.row_dimensions[subtotal_row].outlineLevel = 2
    #     ws.row_dimensions[subtotal_row].hidden = True

    #     start_row += 1  # For the main EBITDA row next
    #     return start_row, subcomponent_rows

    # vector below, utilising C in pandas

    def write_ebitda_subcomponents(self, ws, df, start_row):
        subcomponents = ["Amortization", "Interest expense", "Income taxes"]

        # Extract sales and subcomponents
        sales_series = df.loc[df['label'] == 'Sales'].iloc[0, 1:]
        sub_df = df.loc[df['label'].isin(subcomponents)]

        # Totals & % of sales (vectorized)
        totals = sub_df.iloc[:, 1:].sum()
        pct_of_sales = sub_df.iloc[:, 1:].div(sales_series)

        # Write subcomponents
        for label, row_vals, pct_vals in zip(
            sub_df['label'], sub_df.iloc[:, 1:].values, pct_of_sales.values
        ):
            ws.cell(row=start_row, column=1, value=label).font = self.base.arial_font
            for j, (val, pct) in enumerate(zip(row_vals, pct_vals), start=2):
                ws.cell(row=start_row, column=j*2-2, value=val)  # value
                ws.cell(row=start_row, column=j*2-1, value=pct)  # %
            start_row += 1

        # Subtotal row
        ws.cell(row=start_row, column=1, value="total").font = self.base.bold_font
        for j, (val, pct) in enumerate(zip(totals, totals / sales_series), start=2):
            ws.cell(row=start_row, column=j*2-2, value=val)
            ws.cell(row=start_row, column=j*2-1, value=pct)

        return start_row + 1
